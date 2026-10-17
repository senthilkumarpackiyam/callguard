"""Downloads for the website: the vulnerability list as CSV, Excel or PDF, and a chat answer as PDF."""

import csv
import io
import re
from datetime import datetime, timezone

import markdown
from fpdf import FPDF, FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from callguard.index import ROOT

LOGO = ROOT / "web" / "static" / "brand" / "logo-mark.png"
NVD_PAGE = "https://nvd.nist.gov/vuln/detail/{}"
NVD_NOTICE = "This product uses the NVD API but is not endorsed or certified by the NVD."
SAFETY_NOTE = "No CVE found does not mean a product or version is safe. Check vendor advisories."
NAVY, TEAL, SLATE = (14, 43, 77), (11, 140, 128), (91, 107, 127)
COLUMNS = ["CVE", "Published", "Products", "Vendors", "Severity", "CVSS", "Exploited (CISA KEV)", "Ransomware",
           "EPSS (30-day)", "Description", "NVD page"]


def row(r):
    return [r["id"], r["published"], "; ".join(p["name"] for p in r["products"]),
            "; ".join(sorted({p["vendor"] for p in r["products"]})),
            (r["cvss"] or {}).get("severity") or "", (r["cvss"] or {}).get("score") or "",
            f"since {r['exploited']['since']}" if r["exploited"] else "no",
            (r["exploited"] or {}).get("ransomware") or "",
            f"{r['epss']['score']:.1%}" if r["epss"] else "", r["description"], NVD_PAGE.format(r["id"])]


def as_of_line(manifest):
    return (f"Data as of {manifest['collected_at'][:16].replace('T', ' ')} UTC "
            f"(CISA KEV catalog {manifest['kev']['catalog_version']}, EPSS {manifest['epss']['score_date']})")


def cves_csv(records):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    writer.writerows(row(r) for r in records)
    return out.getvalue().encode("utf-8-sig")  # the signature lets Excel read accented names correctly


def cves_xlsx(records, manifest, filters):
    book = Workbook()
    sheet = book.active
    sheet.title = "CVEs"
    sheet.append(["CallGuard: contact-center vulnerabilities"])
    sheet.append([as_of_line(manifest) + (f" | Filters: {filters}" if filters else "")])
    sheet.append([])
    sheet.append(COLUMNS)
    sheet["A1"].font = Font(bold=True, size=14, color="0E2B4D")
    sheet["A2"].font = Font(italic=True, color="5B6B7F")
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E2B4D")
    for r in records:
        sheet.append(row(r))
        link = sheet.cell(row=sheet.max_row, column=len(COLUMNS))
        link.hyperlink, link.font = link.value, Font(color="16406E", underline="single")
    for column, width in zip("ABCDEFGHIJK", (17, 12, 34, 16, 11, 7, 18, 12, 12, 90, 44)):
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=5, min_col=10, max_col=10):
        cells[0].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "B5"
    sheet.auto_filter.ref = f"A4:K{sheet.max_row}"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ---------- PDF ----------

REPLACEMENTS = {"“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-", "…": "...", "→": "->", "←": "<-",
                "≤": "<=", "≥": ">=", "•": "-", "·": "-", " ": " ", "✓": "yes", "✗": "no"}


def latin1(text):
    """The PDF's built-in fonts only cover Latin-1, so swap common symbols for plain ones."""
    for a, b in REPLACEMENTS.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


class Report(FPDF):
    def __init__(self, title, subtitle, orientation="P"):
        super().__init__(orientation=orientation, format="A4")
        self.report_title, self.subtitle = title, subtitle
        self.set_auto_page_break(auto=True, margin=16)
        self.set_margins(14, 14, 14)
        self.add_page()

    def header(self):
        if LOGO.exists():
            self.image(str(LOGO), x=14, y=10, w=11)
        # text() places words exactly, so "Call" and "Guard" join into one word (cells add padding between them)
        self.set_font("helvetica", "B", 15)
        self.set_text_color(*NAVY)
        self.text(27, 15.6, "Call")
        self.set_text_color(*TEAL)
        self.text(27 + self.get_string_width("Call"), 15.6, "Guard")
        self.set_font("helvetica", "", 8)
        self.set_text_color(*SLATE)
        self.text(27, 20, "CONTACT-CENTER VULNERABILITY INTELLIGENCE")
        self.set_draw_color(226, 232, 240)
        self.line(14, 23, self.w - 14, 23)
        self.set_y(28)

    def footer(self):
        # Two short lines, so the notices never run into the page number
        width = self.w - self.l_margin - self.r_margin
        self.set_y(-13)
        self.set_font("helvetica", "", 7.5)
        self.set_text_color(*SLATE)
        self.cell(width - 24, 3.6, latin1(NVD_NOTICE))
        self.cell(24, 3.6, f"Page {self.page_no()} of {{nb}}", align="R", new_x="LMARGIN", new_y="NEXT")
        self.cell(width - 24, 3.6, latin1(SAFETY_NOTE))

    def title_block(self):
        self.set_font("helvetica", "B", 16)
        self.set_text_color(*NAVY)
        self.multi_cell(0, 8, latin1(self.report_title), new_x="LMARGIN", new_y="NEXT")
        self.set_font("helvetica", "", 9)
        self.set_text_color(*SLATE)
        self.multi_cell(0, 5, latin1(self.subtitle), new_x="LMARGIN", new_y="NEXT")
        self.ln(3)


def cves_pdf(records, manifest, filters):
    pdf = Report(f"Contact-center vulnerabilities ({len(records)})",
                 as_of_line(manifest) + (f"\nFilters: {filters}" if filters else ""), orientation="L")
    pdf.title_block()
    pdf.set_font("helvetica", "", 8.5)
    pdf.set_text_color(27, 42, 58)
    headings = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=NAVY)
    with pdf.table(col_widths=(29, 20, 50, 21, 28, 121), headings_style=headings, line_height=4.6,
                   text_align=("LEFT", "LEFT", "LEFT", "LEFT", "LEFT", "LEFT"), padding=1.4,
                   borders_layout="HORIZONTAL_LINES") as table:
        table.row(["CVE", "Published", "Products", "Severity", "Exploited", "Summary"])
        for r in records:
            cvss = r["cvss"] or {}
            severity = f"{cvss.get('severity', '').title()} {cvss.get('score', '')}".strip() or "-"
            if not r["exploited"]:
                exploited = "No"
            else:
                exploited = "Yes (ransomware)" if r["exploited"]["ransomware"] == "Known" else "Yes"
            summary = r["description"] if len(r["description"]) <= 260 else r["description"][:257].rstrip(" .,;") + "..."
            table.row([latin1(x) for x in (r["id"], r["published"], ", ".join(p["name"] for p in r["products"]),
                                           severity, exploited, summary)])
    return bytes(pdf.output())


def answer_pdf(question, text, as_of):
    pdf = Report("CallGuard answer", f"Data as of {as_of}\nCreated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC")
    pdf.title_block()
    pdf.set_font("helvetica", "B", 11)
    pdf.set_text_color(*NAVY)
    pdf.multi_cell(0, 6, latin1(f"Question: {question}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    body = re.sub(r"\n*_Data as of .*_\s*$", "", text.strip())  # already in the subtitle
    body = re.sub(r"\[(CVE-\d{4}-\d{4,})\](?!\()", r"\1", body)  # [CVE-...] citations read better as plain IDs
    html = markdown.markdown(latin1(body))
    pdf.set_font("helvetica", "", 10)
    pdf.set_text_color(27, 42, 58)
    pdf.write_html(html, font_family="helvetica", li_prefix_color=TEAL,  # teal bullets, not fpdf's default red
                   tag_styles={"a": FontFace(color=(22, 64, 110), emphasis="UNDERLINE")})
    return bytes(pdf.output())
